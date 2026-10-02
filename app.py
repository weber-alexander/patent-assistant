# app.py
# Patent-Assistent (lokal): Navigation links, Bearbeitungsseiten und Dokumentansicht
# Start: python -m streamlit run app.py

import json
import re
from io import BytesIO
from pathlib import Path

import streamlit as st
import ollama
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_COLOR_INDEX

from anspruchs_checker import pruefen
from wissensbasis import (wissens_dateien, index_aufbauen, index_info,
                          suchen, RAG_PROMPT)
from prompts import (SYSTEM_PROMPT, KORREKTUR_PROMPT,
                     CHECKER_KORREKTUR_PROMPT, ABSCHNITTE)

# ======================================================================
# Grundeinstellungen
# ======================================================================

client = ollama.Client(host="http://127.0.0.1:11434")

MODELLE = ["qwen3:4b-instruct", "qwen3:1.7b"]

# temperature: niedrig = sachlich | num_ctx: wie viel Text das Modell "sieht"
KI_OPTIONEN = {"temperature": 0.3, "num_ctx": 8192}

PROJEKT_ORDNER = Path("projekte")
PROJEKT_ORDNER.mkdir(exist_ok=True)

FELDER = ["projektname", "erfindung", "stand_bekannt"] + list(ABSCHNITTE)
BESCHREIBUNGSTEILE = ["gebiet", "stand", "aufgabe", "loesung", "ausfuehrung"]

PLATZHALTER = re.compile(r"(\[ERGÄNZEN:.*?\])", re.DOTALL)
# Typische Länge (Zeichen) je Abschnitt. Nur für die Schätzung des Fortschrittsbalkens.
ERWARTETE_LAENGE = {
    "ansprueche": 1800, "titel": 80, "gebiet": 250, "stand": 600,
    "aufgabe": 250, "loesung": 900, "ausfuehrung": 1800,
    "zusammenfassung": 700, "fragen": 1600,
}

STAND_PLATZHALTER = ("[ERGÄNZEN: Stand der Technik. Bitte zuerst auf der Seite "
                     "'Erfindung' das Feld 'Bekannter Stand der Technik' ausfüllen.]")

# Suchbegriffe für die Wissensbasis pro Abschnitt
RAG_SUCHE = {
    "ansprueche": "Patentansprüche Form unabhängige und abhängige Ansprüche "
                  "Oberbegriff kennzeichnender Teil Rückbeziehung Klarheit",
    "titel": "Bezeichnung der Erfindung kurz und genau technische Bezeichnung",
    "gebiet": "Beschreibung technisches Gebiet der Erfindung angeben",
    "stand": "Beschreibung Stand der Technik angeben",
    "aufgabe": "Beschreibung technische Aufgabe Problem der Erfindung",
    "loesung": "Beschreibung Lösung der Aufgabe vorteilhafte Wirkungen",
    "ausfuehrung": "Ausführung der Erfindung deutlich und vollständig offenbaren "
                   "Fachmann Ausführungsbeispiel",
    "zusammenfassung": "Zusammenfassung der Anmeldung Inhalt Umfang Zeichen",
}

SEITEN = {
    "erfindung": "💡  Erfindung",
    "ansprueche": "⚖️  Patentansprüche",
    "beschreibung": "📝  Beschreibung",
    "fragen": "❓  Erfinderfragen",
    "dokument": "📄  Dokumentansicht",
    "wissen": "📚  Wissensbasis",
}

# ======================================================================
# Seite und Design
# ======================================================================

st.set_page_config(page_title="Patent-Assistent", page_icon="📄", layout="wide")

st.markdown("""
<style>
/* Allgemein */
footer {visibility: hidden;}
.block-container {padding-top: 2rem; max-width: 1100px;}
h1, h2, h3 {color: #2C4A66;}

/* Seitenleiste */
section[data-testid="stSidebar"] {background-color: #E2E7EC;}
.nav-titel {font-size: 0.75rem; font-weight: 600; letter-spacing: 0.08em;
            color: #5B6B7A; margin: 0.8rem 0 0.2rem 0;}

/* Das "Blatt" der Dokumentansicht */
.st-key-papier {
    background: #FFFFFF;
    max-width: 840px;
    margin: 0 auto;
    padding: 56px 72px !important;
    box-shadow: 0 2px 14px rgba(31, 41, 51, 0.14);
    border-radius: 2px;
}
.st-key-papier [data-baseweb="textarea"],
.st-key-papier [data-baseweb="base-input"] {
    border: none !important;
    background: transparent !important;
}
.st-key-papier [data-baseweb="textarea"]:hover {background: #F6F8FA !important;}
.st-key-papier textarea {
    font-family: Arial, sans-serif;
    font-size: 15px;
    line-height: 1.6;
    color: #1F2933;
    padding: 4px 2px !important;
    resize: none;
}
.st-key-papier textarea:focus {background: #F0F4F8 !important;}
.st-key-doc_titel textarea {font-size: 24px !important; font-weight: 700;
                            color: #2C4A66 !important;}
.doc-h1 {font-family: Arial, sans-serif; font-size: 19px; font-weight: 700;
         color: #2C4A66; margin: 18px 0 4px 0;}
.doc-h2 {font-family: Arial, sans-serif; font-size: 15px; font-weight: 700;
         color: #3B5B7A; margin: 12px 0 0 0;}
.doc-umbruch {text-align: center; color: #9AA5B1; font-size: 12px;
              border-top: 1px dashed #C5CDD5; margin: 28px 0 12px 0;
              padding-top: 4px;}
</style>
""", unsafe_allow_html=True)

# ======================================================================
# Zentraler Datenspeicher
# Alle Inhalte liegen in st.session_state.daten. Die Textfelder auf den
# einzelnen Seiten sind nur "Fenster" darauf.
# ======================================================================

if "daten" not in st.session_state:
    st.session_state.daten = {feld: "" for feld in FELDER}
st.session_state.setdefault("backup", {})          # für "Rückgängig"
st.session_state.setdefault("quellen", {})         # gefundene Richtlinien
st.session_state.setdefault("aktuelle_datei", None)

D = st.session_state.daten

if "meldung" in st.session_state:
    st.toast(st.session_state.pop("meldung"))


# ======================================================================
# Hilfsfunktionen
# ======================================================================

def bereinigen(text):
    """Entfernt einen eventuellen Denktext (alles bis </think>)."""
    if "</think>" in text:
        text = text.split("</think>")[-1]
    return text.replace("<think>", "").strip()


def auto_hoehe(text, minimum=80, zeichen_pro_zeile=90):
    """Schätzt die nötige Höhe eines Textfelds aus der Textlänge."""
    zeilen = sum(len(z) // zeichen_pro_zeile + 1 for z in (text or "").split("\n"))
    return max(minimum, min(1400, zeilen * 26 + 30))


def _uebernehmen(feld, widget_key):
    """Schreibt eine Änderung aus einem Textfeld in den Datenspeicher."""
    D[feld] = st.session_state[widget_key]


def textfeld(feld, label, seite, minimum=80, verbergen=False, platzhalter="",
             zeichen_pro_zeile=90):
    """Zeichnet ein Textfeld, das mit dem Datenspeicher verbunden ist."""
    widget_key = f"w_{seite}_{feld}"
    st.session_state[widget_key] = D[feld]
    st.text_area(
        label, key=widget_key,
        height=auto_hoehe(D[feld], minimum, zeichen_pro_zeile),
        on_change=_uebernehmen, args=(feld, widget_key),
        label_visibility="collapsed" if verbergen else "visible",
        placeholder=platzhalter,
    )


def kontext():
    """Grundlage, die die KI zu jeder Aufgabe bekommt."""
    teile = ["Erfindungsbeschreibung:\n" + D["erfindung"]]
    teile.append("Bekannter Stand der Technik:\n"
                 + (D["stand_bekannt"].strip() or "(keine Angaben vorhanden)"))
    if D["ansprueche"].strip():
        teile.append("Bereits formulierte Patentansprüche:\n" + D["ansprueche"])
    return "\n\n".join(teile)


def richtlinien_kontext(feld):
    """Sucht passende Regeln. Kuratierte Regeln (regeln_*.txt) haben Vorrang."""
    if not st.session_state.get("rag_an") or feld not in RAG_SUCHE:
        return "", []
    anfrage = f"[{ABSCHNITTE[feld]['name']}] " + RAG_SUCHE[feld]
    alle = suchen(anfrage, k=15)
    treffer = [t for t in alle if t["quelle"].startswith("regeln")][:4] or alle[:2]
    if not treffer:
        return "", []
    regeln = "\n".join(f"- {t['text']}" for t in treffer)
    return ("\n\nBEACHTE ZUSÄTZLICH DIESE REGELN (nicht zitieren, nur befolgen):\n"
            + regeln), treffer


def ki_text(system_text, user_text, erwartet, fortschritt):
    """Ruft die KI auf und gibt den fertigen Text zurück.
    fortschritt: Funktion, die einen geschätzten Wert zwischen 0 und 1 bekommt."""
    volltext = ""
    stream = client.chat(
        model=st.session_state["modell"],
        messages=[{"role": "system", "content": system_text},
                  {"role": "user", "content": user_text}],
        stream=True,
        options=KI_OPTIONEN,
    )
    for teil in stream:
        volltext += teil["message"]["content"] or ""
        fortschritt(min(0.95, len(volltext) / max(erwartet, 1)))
    fortschritt(1.0)
    return bereinigen(volltext)


def ki_ausfuehren(feld, system_text, user_text, erwartet=None):
    """Einzelner KI-Auftrag mit Fortschrittsbalken. Ergebnis landet im Feld."""
    erwartet = erwartet or ERWARTETE_LAENGE.get(feld, 1000)
    balken = st.progress(0.0, text="KI schreibt ... 0 %")

    def fortschritt(wert):
        balken.progress(wert, text=f"KI schreibt ... {int(wert * 100)} %")

    text = ki_text(system_text, user_text, erwartet, fortschritt)
    st.session_state.backup[feld] = D[feld]
    D[feld] = text
    st.rerun()


def ki_leiste(feld, seite):
    """Buttons unter einem Abschnitt: Erstellen, Korrekturlesen, Rückgängig."""
    info = ABSCHNITTE[feld]
    sp1, sp2, sp3 = st.columns(3)

    if sp1.button("Mit KI erstellen", key=f"gen_{seite}_{feld}",
                  use_container_width=True):
        if not D["erfindung"].strip():
            st.warning("Bitte zuerst auf der Seite 'Erfindung' eine "
                       "Erfindungsbeschreibung eingeben.")
        elif feld == "stand" and not D["stand_bekannt"].strip():
            st.session_state.backup[feld] = D[feld]
            D[feld] = ("[ERGÄNZEN: Stand der Technik. Bitte zuerst auf der Seite "
                       "'Erfindung' das Feld 'Bekannter Stand der Technik' ausfüllen.]")
            st.rerun()
        else:
            zusatz, quellen = richtlinien_kontext(feld)
            st.session_state.quellen[feld] = quellen
            ki_ausfuehren(feld, SYSTEM_PROMPT + "\n\n" + info["prompt"] + zusatz,
                          kontext())

    if sp2.button("Korrekturlesen", key=f"kor_{seite}_{feld}",
                  use_container_width=True):
        if not D[feld].strip():
            st.warning("Das Feld ist leer.")
        else:
            ki_ausfuehren(feld, KORREKTUR_PROMPT, D[feld], erwartet=len(D[feld]))

    if feld in st.session_state.backup:
        if sp3.button("Rückgängig", key=f"undo_{seite}_{feld}",
                      use_container_width=True):
            D[feld] = st.session_state.backup.pop(feld)
            st.rerun()
def gesamtablauf(mit_fragen):
    """Erstellt den kompletten Entwurf in einem Durchlauf (ohne Korrekturlesen)."""
    schritte = (["ansprueche", "checker", "titel"] + BESCHREIBUNGSTEILE
                + ["zusammenfassung"] + (["fragen"] if mit_fragen else []))
    anzahl = len(schritte)
    balken = st.progress(0.0, text="Starte ...")

    for i, schritt in enumerate(schritte):
        name = ("Anspruchs-Checker" if schritt == "checker"
                else ABSCHNITTE[schritt]["name"])

        def fortschritt(wert, i=i, name=name):
            gesamt = (i + wert) / anzahl
            balken.progress(min(gesamt, 1.0),
                            text=f"Schritt {i + 1}/{anzahl}: {name} ... {int(gesamt * 100)} %")

        fortschritt(0.0)

        # Checker: bei Befunden einmal per KI korrigieren,
        # aber nur übernehmen, wenn es danach weniger Befunde sind
        if schritt == "checker":
            befunde = pruefen(D["ansprueche"])
            if befunde:
                liste = "\n".join(f"- Anspruch {b.anspruch}: {b.text}" for b in befunde)
                neu = ki_text(SYSTEM_PROMPT + "\n\n" + CHECKER_KORREKTUR_PROMPT,
                              "Erfindungsbeschreibung:\n" + D["erfindung"]
                              + "\n\nAnspruchssatz:\n" + D["ansprueche"]
                              + "\n\nPrüfbefunde:\n" + liste,
                              len(D["ansprueche"]), fortschritt)
                if len(pruefen(neu)) < len(befunde):
                    D["ansprueche"] = neu
            continue

        st.session_state.backup[schritt] = D[schritt]

        if schritt == "stand" and not D["stand_bekannt"].strip():
            D[schritt] = STAND_PLATZHALTER
            continue

        zusatz, quellen = richtlinien_kontext(schritt)
        st.session_state.quellen[schritt] = quellen
        D[schritt] = ki_text(SYSTEM_PROMPT + "\n\n" + ABSCHNITTE[schritt]["prompt"] + zusatz,
                             kontext(), ERWARTETE_LAENGE[schritt], fortschritt)

    balken.progress(1.0, text="Fertig")
    st.session_state.meldung = "Gesamtentwurf erstellt. Bitte prüfen!"
    st.session_state["_seite_neu"] = "dokument"    # danach zur Dokumentansicht
    st.rerun()

def abschnitt_bearbeiten(feld, seite):
    """Ein kompletter Abschnitt auf einer Bearbeitungsseite."""
    info = ABSCHNITTE[feld]
    st.markdown(f"#### {info['name']}")
    textfeld(feld, info["name"], seite, minimum=min(info["hoehe"], 150),
             verbergen=True)

    if feld == "zusammenfassung" and D[feld].strip():
        zeichen = len(D[feld])
        if zeichen > 1500:
            st.error(f"{zeichen} Zeichen. Erlaubt sind höchstens 1.500 (§ 13 PatV).")
        else:
            st.caption(f"{zeichen} / 1.500 Zeichen (§ 13 PatV)")

    quellen = st.session_state.quellen.get(feld)
    if quellen:
        with st.expander(f"Berücksichtigte Regeln ({len(quellen)})"):
            for t in quellen:
                st.markdown(f"**{t['quelle']}** (Ähnlichkeit {t['score']:.2f})")
                st.caption(t["text"])

    ki_leiste(feld, seite)


def offene_platzhalter():
    return sum(len(PLATZHALTER.findall(D[k])) for k in ABSCHNITTE if k != "fragen")

def text_aus_datei(datei):
    """Extrahiert Text aus einer hochgeladenen Datei (.docx, .pdf, .txt)."""
    endung = Path(datei.name).suffix.lower()
    inhalt = datei.getvalue()

    if endung == ".txt":
        try:
            text = inhalt.decode("utf-8")
        except UnicodeDecodeError:
            text = inhalt.decode("cp1252", errors="ignore")   # ältere Windows-Dateien

    elif endung == ".docx":
        doc = Document(BytesIO(inhalt))
        teile = [p.text for p in doc.paragraphs if p.text.strip()]
        # Erfindungsmeldungen sind oft Formulare mit Tabellen
        for tabelle in doc.tables:
            for zeile in tabelle.rows:
                zellen = []
                for zelle in zeile.cells:
                    wert = zelle.text.strip()
                    if wert and wert not in zellen:      # verbundene Zellen nur einmal
                        zellen.append(wert)
                if zellen:
                    teile.append(" | ".join(zellen))
        text = "\n".join(teile)

    elif endung == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(inhalt))
        text = "\n".join(seite.extract_text() or "" for seite in reader.pages)
        text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)   # Silbentrennung am Zeilenende

    else:
        return ""

    text = re.sub(r"[ \t]+", " ", text)          # doppelte Leerzeichen
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)  # viele Leerzeilen
    return text.strip()

def dateiname(name):
    sauber = re.sub(r"[^\w\-]", "_", name.strip()).strip("_")
    return sauber or "projekt"


# ---------- Projekte ----------

def speichern(name_datei):
    datei = PROJEKT_ORDNER / (name_datei + ".json")
    datei.write_text(json.dumps({f: D[f] for f in FELDER}, ensure_ascii=False,
                                indent=2), encoding="utf-8")
    st.session_state.aktuelle_datei = name_datei
    st.session_state.meldung = f"Gespeichert: {name_datei}.json"
    st.rerun()


def projekt_setzen(daten, datei_name):
    D.clear()
    D.update({f: daten.get(f, "") for f in FELDER})
    st.session_state.backup = {}
    st.session_state.quellen = {}
    st.session_state.aktuelle_datei = datei_name
    st.rerun()


# ---------- Word-Export ----------

def text_einfuegen(doc, text):
    """Fügt Text absatzweise ein und markiert Platzhalter gelb."""
    for zeile in text.strip().split("\n"):
        zeile = zeile.strip()
        if not zeile:
            continue
        absatz = doc.add_paragraph()
        for teil in PLATZHALTER.split(zeile):
            if not teil:
                continue
            run = absatz.add_run(teil)
            if PLATZHALTER.fullmatch(teil):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                run.bold = True


def docx_erstellen():
    doc = Document()
    stil = doc.styles["Normal"]
    stil.font.name = "Arial"
    stil.font.size = Pt(11)

    doc.add_heading(D["titel"].strip() or "[ERGÄNZEN: Titel]", level=0)
    doc.add_heading("Beschreibung", level=1)
    for key in BESCHREIBUNGSTEILE:
        name = ABSCHNITTE[key]["name"]
        doc.add_heading(name, level=2)
        text_einfuegen(doc, D[key] or f"[ERGÄNZEN: {name}]")

    doc.add_page_break()
    doc.add_heading("Patentansprüche", level=1)
    text_einfuegen(doc, D["ansprueche"] or "[ERGÄNZEN: Patentansprüche]")

    doc.add_page_break()
    doc.add_heading("Zusammenfassung", level=1)
    text_einfuegen(doc, D["zusammenfassung"] or "[ERGÄNZEN: Zusammenfassung]")

    puffer = BytesIO()
    doc.save(puffer)
    return puffer.getvalue()


def export_button(key):
    st.download_button(
        "📥 Als Word exportieren",
        data=docx_erstellen(),
        file_name=dateiname(D["projektname"] or "Anmeldung") + ".docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True, type="primary", key=key,
    )


# ======================================================================
# Seitenleiste
# ======================================================================

with st.sidebar:
    st.markdown("## Patent-Assistent")
    st.caption("Lokal · vertraulich · kein Rechtsrat")

    # ----- Projekt -----
    st.markdown('<p class="nav-titel">PROJEKT</p>', unsafe_allow_html=True)
    aktuell = st.session_state.aktuelle_datei
    st.caption(f"Geöffnet: **{aktuell}.json**" if aktuell
               else "Neues, noch nicht gespeichertes Projekt")
    textfeld_key = "w_side_projektname"
    st.session_state[textfeld_key] = D["projektname"]
    st.text_input("Projektname", key=textfeld_key, placeholder="Projektname",
                  label_visibility="collapsed",
                  on_change=_uebernehmen, args=("projektname", textfeld_key))

    sp_a, sp_b = st.columns(2)
    if sp_a.button("💾 Speichern", use_container_width=True):
        speichern(aktuell or dateiname(D["projektname"]))
    if sp_b.button("📝 Speichern unter", use_container_width=True):
        ziel = dateiname(D["projektname"])
        if ziel != aktuell and (PROJEKT_ORDNER / (ziel + ".json")).exists():
            st.error(f"'{ziel}.json' existiert bereits.")
        else:
            speichern(ziel)

    vorhandene = sorted(p.stem for p in PROJEKT_ORDNER.glob("*.json"))
    if vorhandene:
        index = vorhandene.index(aktuell) if aktuell in vorhandene else 0
        auswahl = st.selectbox("Projekt öffnen", vorhandene, index=index,
                               label_visibility="collapsed")
        sp_c, sp_d = st.columns(2)
        if sp_c.button("📂 Laden", use_container_width=True):
            daten = json.loads((PROJEKT_ORDNER / (auswahl + ".json"))
                               .read_text(encoding="utf-8"))
            projekt_setzen(daten, auswahl)
        if sp_d.button("🆕 Neu", use_container_width=True):
            projekt_setzen({}, None)
    elif st.button("🆕 Neues Projekt", use_container_width=True):
        projekt_setzen({}, None)

    # ----- Navigation -----
    st.markdown('<p class="nav-titel">BEARBEITEN</p>', unsafe_allow_html=True)
    if "_seite_neu" in st.session_state:
        st.session_state["seite"] = st.session_state.pop("_seite_neu")
    seite = st.radio("Navigation", list(SEITEN), format_func=SEITEN.get,
                     key="seite", label_visibility="collapsed")

    # ----- Einstellungen -----
    st.markdown('<p class="nav-titel">EINSTELLUNGEN</p>', unsafe_allow_html=True)
    with st.expander("⚙️ KI-Einstellungen"):
        st.selectbox("Modell", MODELLE, key="modell")
        st.toggle("Richtlinien berücksichtigen (RAG)", key="rag_an",
                  help="Gibt der KI beim Erstellen passende Regeln aus der "
                       "Wissensbasis mit. Etwas langsamer.")


# ======================================================================
# Seiten
# ======================================================================

if seite == "erfindung":
    st.title("Erfindung")
    st.caption("Grundlage für alle weiteren Texte. Je genauer, desto besser.")
    st.markdown("#### Erfindungsmeldung hochladen")
    st.caption("Word (.docx), PDF oder Text. Die Datei wird nur lokal verarbeitet "
               "und nicht gespeichert. Gescannte PDFs ohne Textebene und das alte "
               ".doc-Format werden nicht unterstützt.")
    datei = st.file_uploader("Datei auswählen", type=["docx", "pdf", "txt"],
                             label_visibility="collapsed")

    if datei:
        # Eigener Speicherplatz pro Datei: Text wird nur einmal ausgelesen,
        # damit Änderungen in der Vorschau erhalten bleiben
        vorschau_key = f"vorschau_{datei.name}_{datei.size}"
        if vorschau_key not in st.session_state:
            try:
                st.session_state[vorschau_key] = text_aus_datei(datei)
            except Exception as fehler:
                st.session_state[vorschau_key] = ""
                st.error(f"Die Datei konnte nicht gelesen werden: {fehler}")

        if not st.session_state[vorschau_key].strip():
            st.warning("Kein Text gefunden. Handelt es sich um ein gescanntes "
                       "PDF? Dann bitte den Text manuell einfügen.")
        else:
            st.text_area("Vorschau (bearbeitbar, z. B. Erfindernamen oder "
                         "Formularfelder entfernen)",
                         key=vorschau_key, height=250)
            text = st.session_state[vorschau_key]

            st.caption(f"{len(text):,} Zeichen".replace(",", "."))
            if len(text) > 15000:
                st.warning("Sehr langer Text. Das Modell berücksichtigt eventuell "
                           "nicht alles. Am besten auf die technisch relevanten "
                           "Teile kürzen.")

            ziel = st.radio("Übernehmen in", ["erfindung", "stand_bekannt"],
                            format_func={"erfindung": "Erfindungsbeschreibung",
                                         "stand_bekannt": "Bekannter Stand der Technik"}.get,
                            horizontal=True)
            sp_e, sp_a = st.columns(2)
            if sp_e.button("Ersetzen", use_container_width=True, type="primary"):
                D[ziel] = text.strip()
                st.session_state.meldung = "Text übernommen."
                st.rerun()
            if sp_a.button("Anhängen", use_container_width=True):
                D[ziel] = (D[ziel].rstrip() + "\n\n" + text.strip()).strip()
                st.session_state.meldung = "Text angehängt."
                st.rerun()

    st.divider()
    st.markdown("#### Erfindungsbeschreibung")
    textfeld("erfindung", "Erfindungsbeschreibung", seite, minimum=250,
             verbergen=True,
             platzhalter="Merkmale, Funktionsweise, Zahlenwerte, Varianten, Vorteile ...")
    st.markdown("#### Bekannter Stand der Technik")
    textfeld("stand_bekannt", "Bekannter Stand der Technik", seite, minimum=120,
             verbergen=True,
             platzhalter="Was gibt es bisher, und was sind die Nachteile? "
                         "Leer lassen, wenn unbekannt.")
    
    st.divider()
    st.markdown("#### Gesamtentwurf")
    st.caption("Erstellt nacheinander Ansprüche, Checker-Korrektur, Titel, "
               "Beschreibung und Zusammenfassung. Korrekturlesen bitte danach "
               "gezielt pro Abschnitt. Dauer: einige Minuten.")
    mit_fragen = st.checkbox("Erfinderfragen mit erstellen", value=True)
    if any(D[k].strip() for k in ABSCHNITTE):
        st.warning("Vorhandene Texte werden überschrieben. Jeder Abschnitt lässt "
                   "sich danach einzeln per 'Rückgängig' wiederherstellen.")
    if st.button("Gesamten Entwurf erstellen", type="primary",
                 use_container_width=True, disabled=not D["erfindung"].strip()):
        gesamtablauf(mit_fragen)


elif seite == "ansprueche":
    st.title("Patentansprüche")
    st.caption("Zuerst die Ansprüche erstellen, denn die Beschreibung baut darauf auf.")
    abschnitt_bearbeiten("ansprueche", seite)

    st.divider()
    st.markdown("#### Anspruchs-Checker")
    st.caption("Regelbasierte Prüfung ohne KI. Fehler sind formale Mängel, "
               "Warnungen sollten fachlich bewertet werden.")

    befunde = pruefen(D["ansprueche"])

    if "befunde_vorher" in st.session_state:
        vorher = st.session_state.pop("befunde_vorher")
        if len(befunde) < vorher:
            st.success(f"KI-Korrektur: {vorher} → {len(befunde)} Befunde ✅")
        else:
            st.error(f"KI-Korrektur: {vorher} → {len(befunde)} Befunde. "
                     "Keine Verbesserung, eventuell 'Rückgängig' klicken.")

    if not D["ansprueche"].strip():
        st.info("Noch keine Ansprüche vorhanden.")
    elif not befunde:
        st.success("Keine Auffälligkeiten gefunden.")
    else:
        anzahl = {s: sum(b.stufe == s for b in befunde)
                  for s in ["Fehler", "Warnung", "Hinweis"]}
        st.write(f"{anzahl['Fehler']} Fehler · {anzahl['Warnung']} Warnungen · "
                 f"{anzahl['Hinweis']} Hinweise")
        for b in befunde:
            ort = f"**Anspruch {b.anspruch}:** " if b.anspruch else "**Gesamt:** "
            {"Fehler": st.error, "Warnung": st.warning}.get(b.stufe, st.info)(ort + b.text)

        if st.button("Befunde mit KI beheben", use_container_width=True):
            liste = "\n".join(f"- Anspruch {b.anspruch}: {b.text}" for b in befunde)
            st.session_state.befunde_vorher = len(befunde)
            ki_ausfuehren("ansprueche",
                          SYSTEM_PROMPT + "\n\n" + CHECKER_KORREKTUR_PROMPT,
                          "Erfindungsbeschreibung:\n" + D["erfindung"]
                          + "\n\nAnspruchssatz:\n" + D["ansprueche"]
                          + "\n\nPrüfbefunde:\n" + liste,
                          erwartet=len(D["ansprueche"]))


elif seite == "beschreibung":
    st.title("Beschreibung")
    if not D["ansprueche"].strip():
        st.warning("Noch keine Ansprüche vorhanden. Die Begriffe in der "
                   "Beschreibung sind dann eventuell uneinheitlich.")
    for feld in ["titel"] + BESCHREIBUNGSTEILE + ["zusammenfassung"]:
        abschnitt_bearbeiten(feld, seite)
        st.divider()


elif seite == "fragen":
    st.title("Fragen an den Erfinder")
    st.caption("Interne Arbeitsnotiz. Wird nicht exportiert.")
    abschnitt_bearbeiten("fragen", seite)


elif seite == "wissen":
    st.title("Wissensbasis")
    st.caption("Regeln und Richtlinien aus dem Ordner 'wissen'. Werden beim "
               "Erstellen genutzt, wenn RAG in den Einstellungen aktiv ist.")
    dateien = wissens_dateien()
    if dateien:
        st.write("**Dokumente:** " + ", ".join(d.name for d in dateien))
    else:
        st.warning("Der Ordner 'wissen' ist leer.")
    st.write(index_info())

    if st.button("Index aufbauen / aktualisieren", disabled=not dateien):
        balken = st.progress(0.0, text="Erstelle Index ...")
        anzahl = index_aufbauen(
            lambda w: balken.progress(w, text=f"Erstelle Index ... {int(w * 100)} %"))
        st.session_state.meldung = f"Index erstellt: {anzahl} Textabschnitte."
        st.rerun()

    st.divider()
    st.markdown("#### Nachschlagen")
    frage = st.text_input("Frage", placeholder="z. B. Wie lang darf die "
                          "Zusammenfassung sein?")
    if st.button("Suchen", disabled=not frage.strip()):
        treffer = suchen(frage, k=4)
        if not treffer:
            st.error("Kein Index vorhanden.")
        else:
            auszuege = "\n\n".join(f"[{i}] ({t['quelle']}) {t['text']}"
                                   for i, t in enumerate(treffer, 1))
            balken = st.progress(0.0, text="KI antwortet ... 0 %")
            antwort = ki_text(
                RAG_PROMPT, f"Auszüge:\n{auszuege}\n\nFrage: {frage}", 600,
                lambda w: balken.progress(w, text=f"KI antwortet ... {int(w * 100)} %"))
            balken.empty()
            st.markdown(antwort)
            with st.expander("Quellen"):
                for i, t in enumerate(treffer, 1):
                    st.markdown(f"**[{i}] {t['quelle']}** ({t['score']:.2f})")
                    st.caption(t["text"])


elif seite == "dokument":
    # ----- Werkzeugleiste über dem Blatt -----
    befunde = pruefen(D["ansprueche"])
    fehler = sum(b.stufe == "Fehler" for b in befunde)
    warnungen = sum(b.stufe == "Warnung" for b in befunde)
    offen = offene_platzhalter()

    sp1, sp2 = st.columns([3, 1])
    with sp1:
        st.markdown("### Dokumentansicht")
        status = [f"{offen} offene Platzhalter",
                  f"Checker: {fehler} Fehler, {warnungen} Warnungen"]
        st.caption(" · ".join(status) + " · Klicke in einen Text, um ihn zu bearbeiten.")
    with sp2:
        st.write("")
        export_button("export_dokument")

    # ----- Das Blatt -----
    with st.container(key="papier"):
        with st.container(key="doc_titel"):
            textfeld("titel", "Titel", seite, minimum=68, verbergen=True,
                     platzhalter="[Titel der Erfindung]", zeichen_pro_zeile=50)

        st.markdown('<div class="doc-h1">Beschreibung</div>', unsafe_allow_html=True)
        for feld in BESCHREIBUNGSTEILE:
            name = ABSCHNITTE[feld]["name"]
            st.markdown(f'<div class="doc-h2">{name}</div>', unsafe_allow_html=True)
            textfeld(feld, name, seite, minimum=68, verbergen=True,
                     platzhalter=f"[{name} noch leer]", zeichen_pro_zeile=85)

        st.markdown('<div class="doc-umbruch">Seitenumbruch</div>',
                    unsafe_allow_html=True)
        st.markdown('<div class="doc-h1">Patentansprüche</div>', unsafe_allow_html=True)
        textfeld("ansprueche", "Patentansprüche", seite, minimum=68, verbergen=True,
                 platzhalter="[Patentansprüche noch leer]", zeichen_pro_zeile=85)

        st.markdown('<div class="doc-umbruch">Seitenumbruch</div>',
                    unsafe_allow_html=True)
        st.markdown('<div class="doc-h1">Zusammenfassung</div>', unsafe_allow_html=True)
        textfeld("zusammenfassung", "Zusammenfassung", seite, minimum=68,
                 verbergen=True, platzhalter="[Zusammenfassung noch leer]",
                 zeichen_pro_zeile=85)

    st.caption("Platzhalter [ERGÄNZEN: ...] werden im Word-Dokument gelb markiert. "
               "Formatierungen wie Fett oder Zeichnungen bitte in Word ergänzen.")