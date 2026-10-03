"""Tests for the Word export."""

from io import BytesIO

from docx import Document

from patent_assistant import projects
from patent_assistant.document_export import build_docx, count_placeholders


def test_placeholders_in_questions_are_not_counted():
    data = projects.empty_project() | {
        "claims": "1. [ERGÄNZEN: Merkmal]",
        "embodiment": "[ERGÄNZEN: Wert] und [ERGÄNZEN: Material]",
        "questions": "[ERGÄNZEN: intern]",
    }
    assert count_placeholders(data) == 3


def test_export_has_application_structure():
    data = projects.empty_project() | {"title": "Messsystem", "claims": "1. Messsystem."}
    document = Document(BytesIO(build_docx(data)))
    headings = [
        p.text for p in document.paragraphs if p.style.name.startswith(("Heading", "Title"))
    ]
    assert headings[0] == "Messsystem"
    assert "Beschreibung" in headings
    assert (
        headings.index("Beschreibung")
        < headings.index("Patentansprüche")
        < headings.index("Zusammenfassung")
    )


def test_empty_sections_become_highlighted_placeholders():
    document = Document(BytesIO(build_docx(projects.empty_project())))
    highlighted = [
        run.text for p in document.paragraphs for run in p.runs if run.font.highlight_color
    ]
    assert "[ERGÄNZEN: Patentansprüche]" in highlighted
