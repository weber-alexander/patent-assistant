"""Tests for extracting text from uploaded files."""

from io import BytesIO

import pytest
from docx import Document

from patent_assistant.file_import import extract_text


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Ein Mähroboter mit einer Wärmebildkamera.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Erfinder"
    table.rows[0].cells[1].text = "Max Muster"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_txt_utf8():
    assert extract_text("a.txt", "Größe".encode()) == "Größe"


def test_txt_legacy_windows_encoding():
    assert extract_text("a.txt", "Größe".encode("cp1252")) == "Größe"


def test_docx_paragraphs_and_tables():
    text = extract_text("meldung.docx", make_docx())
    assert "Wärmebildkamera" in text
    assert "Erfinder | Max Muster" in text


def test_whitespace_is_cleaned():
    assert extract_text("a.txt", b"a   b\n\n\n\n\nc") == "a b\n\nc"


def test_unsupported_format_raises():
    with pytest.raises(ValueError):
        extract_text("alt.doc", b"")
