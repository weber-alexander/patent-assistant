"""Word export in the order of a German patent application."""

from __future__ import annotations

import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt

from patent_assistant.prompts import DESCRIPTION_PARTS, SECTIONS

PLACEHOLDER = re.compile(r"(\[ERGÄNZEN:.*?\])", re.DOTALL)


def count_placeholders(data: dict[str, str]) -> int:
    """Open placeholders in all exported sections (questions are internal)."""
    return sum(
        len(PLACEHOLDER.findall(data.get(key, ""))) for key in SECTIONS if key != "questions"
    )


def _add_text(document: Document, text: str) -> None:
    """Add text paragraph by paragraph and highlight placeholders."""
    for line in text.strip().split("\n"):
        line = line.strip()
        if not line:
            continue
        paragraph = document.add_paragraph()
        for part in PLACEHOLDER.split(line):
            if not part:
                continue
            run = paragraph.add_run(part)
            if PLACEHOLDER.fullmatch(part):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                run.bold = True


def _section_text(data: dict[str, str], key: str) -> str:
    return data.get(key, "").strip() or f"[ERGÄNZEN: {SECTIONS[key]['label']}]"


def build_docx(data: dict[str, str]) -> bytes:
    """Build the application as a .docx file and return its bytes."""
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    document.add_heading(_section_text(data, "title"), level=0)
    document.add_heading("Beschreibung", level=1)
    for key in DESCRIPTION_PARTS:
        document.add_heading(SECTIONS[key]["label"], level=2)
        _add_text(document, _section_text(data, key))

    for key in ("claims", "abstract"):
        document.add_page_break()
        document.add_heading(SECTIONS[key]["label"], level=1)
        _add_text(document, _section_text(data, key))

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
