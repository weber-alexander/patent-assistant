"""Extract plain text from uploaded invention disclosures."""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED_TYPES = ("docx", "pdf", "txt")


def extract_text(filename: str, data: bytes) -> str:
    """Return the text of a .docx, .pdf or .txt file."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        text = _from_txt(data)
    elif suffix == ".docx":
        text = _from_docx(data)
    elif suffix == ".pdf":
        text = _from_pdf(data)
    else:
        raise ValueError(f"Nicht unterstütztes Dateiformat: {suffix}")
    return _clean(text)


def _from_txt(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="ignore")  # legacy Windows files


def _from_docx(data: bytes) -> str:
    document = Document(BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    # Invention disclosures are often forms with tables
    for table in document.tables:
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                value = cell.text.strip()
                if value and value not in cells:  # merged cells appear repeatedly
                    cells.append(value)
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _from_pdf(data: bytes) -> str:
    text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(data)).pages)
    return re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # undo hyphenation at line ends


def _clean(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()
